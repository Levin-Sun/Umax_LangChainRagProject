# TDD 红灯：解析器注册表（模块化可替换，§3.2-③）与扫描件路由 MinerU
import io

import pytest

from app.services.parsers import (
    ScannedPdfError, parse_document, supported_extensions,
)


def test_txt_md_passthrough():
    assert parse_document("a.txt", "你好\n世界".encode()) == "你好\n世界"
    assert parse_document("b.md", "# 标题\n正文".encode()).startswith("# 标题")


def test_unsupported_extension_raises():
    with pytest.raises(ValueError):
        parse_document("x.exe", b"MZ")


def test_docx_paragraphs_and_tables():
    import docx

    d = docx.Document()
    d.add_paragraph("文档标题段")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "SKU"
    t.cell(0, 1).text = "价格"
    t.cell(1, 0).text = "A001"
    t.cell(1, 1).text = "29.9"
    buf = io.BytesIO()
    d.save(buf)

    out = parse_document("手册.docx", buf.getvalue())
    assert "文档标题段" in out
    assert "SKU" in out and "29.9" in out


def test_xlsx_rows_as_text():
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "价格表"
    ws.append(["SKU", "结算账期"])
    ws.append(["A001", "T+7"])
    buf = io.BytesIO()
    wb.save(buf)

    out = parse_document("价格.xlsx", buf.getvalue())
    assert "价格表" in out and "结算账期" in out and "T+7" in out


def test_pptx_slide_texts():
    from pptx import Presentation

    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[5])
    s.shapes.title.text = "618大促规则"
    buf = io.BytesIO()
    p.save(buf)

    assert "618大促规则" in parse_document("培训.pptx", buf.getvalue())


def test_pdf_text_layer(monkeypatch):
    """pypdf 打桩：文字层可抽则抽；无文字层页占比过半判扫描件。"""
    import app.services.parsers as parsers

    class FakePage:
        def __init__(self, text):
            self._t = text

        def extract_text(self):
            return self._t

    holder: list[str] = ["第一页内容"]

    class FakeReader:
        def __init__(self, data):  # 与真 pypdf 一致：入参为字节流
            self.pages = [FakePage(t) for t in holder]

    monkeypatch.setattr(parsers, "_pypdf_reader", FakeReader)
    assert "第一页内容" in parse_document("digital.pdf", b"%PDF-1.4 fake")

    holder[0] = ""  # 全空页=扫描件；未配置 MinerU 时报错并给出启用指引
    with pytest.raises(ScannedPdfError, match="MinerU"):
        parse_document("scan.pdf", b"%PDF-1.4 fake")


def test_scanned_pdf_routes_to_mineru_when_configured(monkeypatch):
    import httpx

    import app.services.parsers as parsers

    class FakePage:
        def extract_text(self):
            return ""

    class FakeReader:
        def __init__(self, data):
            self.pages = [FakePage()]

    monkeypatch.setattr(parsers, "_pypdf_reader", FakeReader)

    def handler(request: httpx.Request) -> httpx.Response:
        assert request.url.path.endswith("/file_parse")
        return httpx.Response(200, json={"results": {"scan.pdf": {"md_content": "OCR还原的扫描合同文字"}}})

    mineru = parsers.MinerUClient(base_url="http://mineru:8002",
                                  transport=httpx.MockTransport(handler))
    out = parse_document("scan.pdf", b"%PDF fake scanned", mineru=mineru)
    assert "OCR还原的扫描合同文字" in out


def test_registry_extends_support():
    assert {".txt", ".md", ".docx", ".xlsx", ".pptx", ".pdf"} <= supported_extensions()
