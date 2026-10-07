# 图表入库（§A「文档里的图表、照片，用视觉模型转成文字描述一起进索引，能被搜到」）：
# 解析层抽出内嵌图片（docx/pptx/xlsx 的 zip media；PDF 的图片 XObject）→ 过滤图标类小图 →
# 视觉模型转文字描述 → 作为 chunk 进索引（meta 标记来源，便于分块预览排查）。
# 视觉失败绝不让整篇文档失败（图只是增强，正文才是主体）；无视觉模型则静默跳过。
import io
import zlib

import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.main import create_app
from app.models import Chunk
from app.services.media import MAX_IMAGES, extract_images
from tests.conftest import login, seed_user

ADMIN = ("admin@umax.local", "Adm1n-Pass-123")


def png_bytes(w: int = 400, h: int = 300, seed: int = 7) -> bytes:
    """带噪声的图：纯色 PNG 压缩后仅 1~2KB，会被"小图=图标"的字节闸滤掉——
    测试夹具要贴近真实的图表/照片（几百 KB 级）。"""
    import random

    from PIL import Image

    rnd = random.Random(seed)
    im = Image.new("RGB", (w, h))
    im.putdata([(rnd.randrange(256), rnd.randrange(256), rnd.randrange(256))
                for _ in range(w * h)])
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    return buf.getvalue()


def docx_with_image(w: int = 400, h: int = 300, extra_images: int = 0) -> bytes:
    import docx

    d = docx.Document()
    d.add_paragraph("售后规则：生鲜商品不支持七天无理由退货。")
    # 每张图内容不同：python-docx 会按内容去重，相同图片只会存一个 part（夹具要防这点）
    for i in range(extra_images + 1):
        d.add_picture(io.BytesIO(png_bytes(w, h, seed=i + 1)))
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def pptx_with_image() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[5])
    slide.shapes.title.text = "退款流程图"
    slide.shapes.add_picture(io.BytesIO(png_bytes()), Inches(1), Inches(1))
    buf = io.BytesIO()
    p.save(buf)
    return buf.getvalue()


def xlsx_with_image() -> bytes:
    import openpyxl
    from openpyxl.drawing.image import Image as XlImage
    from PIL import Image as PILImage

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "库存表"
    ws.add_image(XlImage(PILImage.open(io.BytesIO(png_bytes()))), "B2")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def pdf_with_image(w: int = 300, h: int = 200) -> bytes:
    """手工构造带图片 XObject 的最小 PDF（FlateDecode 原始 RGB，免去 JPEG 素材依赖）。"""
    import random

    rnd = random.Random(11)
    pixels = zlib.compress(bytes(rnd.randrange(256) for _ in range(w * h * 3)))
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        # 带文字层（Resources 里给字体）：纯图无字的 PDF 会被判扫描件而走 MinerU 分支
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 400 400] "
        b"/Resources << /XObject << /Im0 4 0 R >> "
        b"/Font << /F1 6 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /XObject /Subtype /Image /Width " + str(w).encode()
        + b" /Height " + str(h).encode()
        + b" /ColorSpace /DeviceRGB /BitsPerComponent 8 /Filter /FlateDecode /Length "
        + str(len(pixels)).encode() + b" >>\nstream\n" + pixels + b"\nendstream",
        b"<< /Length 96 >>\nstream\nq 400 0 0 400 0 0 cm /Im0 Do Q\n"
        b"BT /F1 14 Tf 20 380 Td (refund policy: within 7 days) Tj ET\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += str(i).encode() + b" 0 obj\n" + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 " + str(len(objs) + 1).encode() + b"\n0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (b"trailer\n<< /Size " + str(len(objs) + 1).encode() + b" /Root 1 0 R >>\n"
            b"startxref\n" + str(xref).encode() + b"\n%%EOF\n")
    return bytes(out)


# ---------------- 抽取层 ----------------

@pytest.mark.parametrize("builder,name", [
    (docx_with_image, "图.docx"), (pptx_with_image, "图.pptx"), (xlsx_with_image, "图.xlsx"),
    (pdf_with_image, "图.pdf"),
])
def test_extract_images_from_all_supported_formats(builder, name):
    blobs = extract_images(name, builder())
    assert len(blobs) == 1, f"{name} 应抽出 1 张图"
    b = blobs[0]
    assert b.data and b.width >= 200 and b.height >= 200
    assert b.ext in {"png", "jpeg", "jpg"}
    assert b.origin  # 记录来源（如 word/media/image1.png）便于排查


def test_small_images_are_filtered_as_icons():
    """小图（图标/项目符号/logo）不该进索引——否则每个项目符号都变成一条描述。"""
    assert extract_images("小.docx", docx_with_image(w=32, h=32)) == []


def test_text_only_documents_have_no_images():
    assert extract_images("a.md", "纯文本内容".encode()) == []


def test_max_images_cap():
    blobs = extract_images("很多.docx", docx_with_image(extra_images=MAX_IMAGES + 5))
    assert len(blobs) == MAX_IMAGES


def test_corrupt_office_file_does_not_raise():
    assert extract_images("坏.docx", b"not a zip at all") == []


# ---------------- 入库链路 ----------------

_DEFAULT_VISION = object()   # 哨兵：显式传 None（"未配视觉模型"）与"不给参数"必须可区分


def _client(engine, tmp_path, vision=_DEFAULT_VISION):
    def chat(q, h):
        return {"answer": "答[1]", "prompt_tokens": 1, "completion_tokens": 1}

    def fake_vision(data_url: str) -> dict:
        assert data_url.startswith("data:image/")
        return {"caption": "退款流程图：含 7 个自然日时限", "prompt_tokens": 50,
                "completion_tokens": 20, "model": "fake-vision"}

    return TestClient(create_app(engine=engine, secret="m-secret", embedder=None,
                                 chat_fn=chat,
                                 vision_fn=fake_vision if vision is _DEFAULT_VISION else vision,
                                 upload_dir=str(tmp_path)))


@pytest.fixture
def client(engine, db, tmp_path):
    c = _client(engine, tmp_path)
    seed_user(engine, *ADMIN, role="admin")
    login(c, *ADMIN)
    c.post("/api/v1/kb", json={"name": "图库"})
    return c


def _upload_and_chunks(c, name: str, raw: bytes, kb=1):
    doc = c.post(f"/api/v1/kb/{kb}/documents",
                 files={"file": (name, raw, "application/octet-stream")}).json()
    chunks = c.get(f"/api/v1/documents/{doc['id']}/chunks").json()
    return doc, chunks


def test_docx_image_caption_enters_index_with_meta(client, engine):
    doc, chunks = _upload_and_chunks(client, "流程.docx", docx_with_image())
    assert doc["status"] == "ready"
    captions = [c for c in chunks if c["meta"].get("source") == "image"]
    assert len(captions) == 1
    assert "退款流程图" in captions[0]["content"]
    assert captions[0]["meta"]["image_index"] == 1
    # 视觉调用记账（场景 vision，供用量看板）
    with Session(engine) as s:
        from app.models import UsageRecord
        rows = s.query(UsageRecord).filter_by(scenario="vision").all()
    assert len(rows) == 1 and rows[0].model == "fake-vision"


def test_vision_failure_keeps_document_ready(client, engine, tmp_path):
    def boom(_url):
        raise RuntimeError("视觉服务挂了")

    c = _client(engine, tmp_path, vision=boom)
    login(c, *ADMIN)
    c.post("/api/v1/kb", json={"name": "库"})
    doc, chunks = _upload_and_chunks(c, "流程.docx", docx_with_image())
    assert doc["status"] == "ready", "图注失败不能把整篇文档判失败"
    assert [c for c in chunks if c["meta"].get("source") == "image"] == []
    assert any("生鲜" in c["content"] for c in chunks)      # 正文照常入库
    assert doc["error"] is None


def test_without_vision_model_no_caption_but_text_ok(engine, db, tmp_path):
    c = _client(engine, tmp_path, vision=None)
    seed_user(engine, *ADMIN, role="admin")
    login(c, *ADMIN)
    c.post("/api/v1/kb", json={"name": "库"})
    doc, chunks = _upload_and_chunks(c, "流程.docx", docx_with_image())
    assert doc["status"] == "ready"
    assert [x for x in chunks if x["meta"].get("source") == "image"] == []
    assert any("生鲜" in x["content"] for x in chunks)


def test_config_center_can_disable_doc_image_caption(client):
    client.put("/api/v1/settings", json={"doc_image_caption": False})
    doc, chunks = _upload_and_chunks(client, "流程.docx", docx_with_image())
    assert doc["status"] == "ready"
    assert [c for c in chunks if c["meta"].get("source") == "image"] == []
    assert client.get("/api/v1/settings").json()["values"]["doc_image_caption"] is False


def test_chunks_preview_exposes_meta_for_debugging(client):
    _, chunks = _upload_and_chunks(client, "流程.docx", docx_with_image())
    assert all("meta" in c for c in chunks)          # §A 分块预览是排查第一工具


def test_reprocess_regenerates_image_captions(client):
    doc, _ = _upload_and_chunks(client, "流程.docx", docx_with_image())
    assert client.post(f"/api/v1/documents/{doc['id']}/reprocess").status_code == 202
    chunks = client.get(f"/api/v1/documents/{doc['id']}/chunks").json()
    assert len([c for c in chunks if c["meta"].get("source") == "image"]) == 1


def test_pdf_image_captioned_too(client):
    doc, chunks = _upload_and_chunks(client, "扫描图.pdf", pdf_with_image())
    assert doc["status"] == "ready", doc["error"]
    assert len([c for c in chunks if c["meta"].get("source") == "image"]) == 1
